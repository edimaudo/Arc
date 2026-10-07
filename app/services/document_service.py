from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {'.pdf', '.xlsx', '.xlsm', '.csv'}


def parse_document(filename: str, data: bytes) -> dict[str, Any]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f'Unsupported file type: {ext or "unknown"}. Use PDF, XLSX, XLSM, or CSV.')
    if ext == '.pdf':
        reader = PdfReader(io.BytesIO(data))
        pages: list[dict[str, Any]] = []
        for i, page in enumerate(reader.pages[:50]):
            text = (page.extract_text() or '').strip()
            if text:
                pages.append({'page': i + 1, 'text': text[:15000]})
        full = '\n\n'.join(f'[Page {x["page"]}]\n{x["text"]}' for x in pages)
        return {'filename': filename, 'type': 'pdf', 'pages': pages, 'text': full[:100000], 'metrics': extract_metrics(full)}
    if ext in {'.xlsx', '.xlsm'}:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets: dict[str, list[list[str]]] = {}
        for ws in wb.worksheets[:20]:
            rows: list[list[str]] = []
            for row in ws.iter_rows(max_row=500, values_only=True):
                vals = [str(v) if v is not None else '' for v in row]
                if any(vals):
                    rows.append(vals)
            sheets[ws.title] = rows
        text = '\n\n'.join(f'[{name}]\n' + '\n'.join(' | '.join(r) for r in rows) for name, rows in sheets.items())
        return {'filename': filename, 'type': 'xlsx', 'sheets': sheets, 'text': text[:100000], 'metrics': extract_metrics(text)}
    content = data.decode('utf-8-sig', errors='replace')
    rows = list(csv.reader(io.StringIO(content)))[:1000]
    text = '\n'.join(' | '.join(r) for r in rows)
    return {'filename': filename, 'type': 'csv', 'rows': rows, 'text': text[:100000], 'metrics': extract_metrics(text)}


def extract_metrics(text: str) -> dict[str, float]:
    import re
    patterns = {
        'revenue_growth': r'revenue\s*(?:growth|increase|change)[^0-9\-+]{0,30}([+\-]?\d+(?:\.\d+)?)\s*%',
        'ebitda_margin': r'ebitda\s*margin[^0-9\-+]{0,30}([+\-]?\d+(?:\.\d+)?)\s*%',
        'net_debt_to_ebitda': r'net\s*debt\s*(?:to|/|:)?\s*ebitda[^0-9]{0,30}([0-9]+(?:\.\d+)?)\s*x?',
        'customer_concentration': r'customer\s*concentration[^0-9]{0,30}(\d+(?:\.\d+)?)\s*%',
    }
    found: dict[str, float] = {}
    for key, pattern in patterns.items():
        m = re.search(pattern, text, re.I)
        if m:
            value = float(m.group(1))
            found[key] = value / 100 if key in {'revenue_growth', 'ebitda_margin'} else value
    return found
